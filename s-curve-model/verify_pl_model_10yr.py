"""
Verify the columns filled by fill_pl_model_10yr.py.

Each new formula must be the previous column's formula with its relative column
references shifted right by one - i.e. the same recursion, next year. Also
checks nothing is left blank and number formats carried over.

    python verify_pl_model_10yr.py [filled.xlsx]
"""
import re
import sys
import openpyxl
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI

OUT = sys.argv[1] if len(sys.argv) > 1 else "UMG2_Video_PL_Model_v135MK_10yr.xlsx"
ws = openpyxl.load_workbook(OUT, data_only=False).active

REF = re.compile(r'(\$?)([A-Z]{1,2})(\$?)(\d{1,4})')

def shift(formula, by=1):
    """Shift relative column references right by `by`. Absolute ($C) stay put."""
    def repl(m):
        dollar_c, col, dollar_r, row = m.groups()
        if dollar_c:                      # absolute column -> unchanged
            return m.group(0)
        return f"{dollar_c}{CL(CI(col) + by)}{dollar_r}{row}"
    return REF.sub(repl, formula)

# Ranges deliberately widened in the new columns (originals stopped at J).
WIDENED = [("$D$81:$J$85", "$D$81:$O$85"), ("$D$80:$J$80", "$D$80:$O$80")]

def normalize(f):
    for old, new in WIDENED:
        f = f.replace(old, new)
    return f

country_rows = [r for b in (19, 29, 39, 49, 59) for r in range(b, b + 7)]
checks = [(country_rows, range(12, 16), 11)]   # L:O continue from K
checks += [(list(range(9, 15)), range(11, 16), 10)]  # K:O continue from J

bad = 0
total = 0
for rows, cols, last_filled in checks:
    for r in rows:
        for c in cols:
            total += 1
            prev = ws.cell(r, c - 1).value
            got  = ws.cell(r, c).value
            assert isinstance(got, str) and got.startswith("="), f"{CL(c)}{r} not a formula"
            if not (isinstance(prev, str) and prev.startswith("=")):
                continue                    # nothing to compare the very first new col to
            want = normalize(shift(prev))
            if normalize(got) != want:
                bad += 1
                print(f"  MISMATCH {CL(c)}{r}\n    got  {got}\n    want {want}")

print(f"checked {total} new cells, {bad} mismatches")
assert bad == 0

# every targeted cell is now populated
empty = [f"{CL(c)}{r}" for rows, cols, _ in checks for r in rows for c in cols
         if ws.cell(r, c).value is None]
print("empty cells remaining in the filled range:", empty or "none")
assert not empty

# number formats carried over from the last filled column
fmt_bad = [f"{CL(c)}{r}" for rows, cols, lf in checks for r in rows for c in cols
           if ws.cell(r, c).number_format != ws.cell(r, lf).number_format]
print("cells with a number format differing from the column they continue:", fmt_bad or "none")
